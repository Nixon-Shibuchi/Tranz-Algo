"""Build Tranz-Algo demo shipment data Word document."""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

OUT = Path(__file__).resolve().parent / "Tranz-Algo_Demo_Shipment_Data.docx"


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)
    doc.add_paragraph("")


def main():
    doc = Document()
    title = doc.add_heading("Tranz-Algo — Demo Shipment Data", level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph(
        "Made-up sample batches for entering into the Tranz-Algo app. "
        "Create products first under Products, then each shipment under Shipments. "
        "Use True Cost to allocate, price, and commit."
    )

    # Shipment 1
    doc.add_heading("Shipment 1 — Already-Local (agent paperwork in ZMW)", level=1)
    doc.add_paragraph(
        "Story: Soaps arrived from Johannesburg. Your clearing agent sent one ZMW invoice — "
        "no foreign currency on the supplier side."
    )

    doc.add_heading("Products (create once)", level=2)
    add_table(
        doc,
        ["SKU", "Name", "Weight (kg)"],
        [
            ["SOAP-LAV", "Lavender bath soap", "0.15"],
            ["SOAP-ALO", "Aloe vera soap", "0.15"],
            ["SOAP-CIT", "Citrus hand soap", "0.12"],
        ],
    )

    doc.add_heading("New shipment header", level=2)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Reference", "SHP-LOC-2026-01"],
            ["Supplier", "Cape Clean Supplies"],
            ["Invoice #", "CA-8842"],
            ["Clearing agent", "Mwamba Logistics"],
            ["Carrier", "Intercity Freight"],
            ["Origin", "South Africa"],
            ["Primary cost mode", "Already-Local"],
            ["Local currency", "ZMW"],
        ],
    )

    doc.add_heading("Lines (mode: Already-Local on each line)", level=2)
    add_table(
        doc,
        ["Product", "Qty", "Local unit (ZMW)"],
        [
            ["SOAP-LAV", "24", "18.50"],
            ["SOAP-ALO", "24", "22.00"],
            ["SOAP-CIT", "12", "19.00"],
        ],
    )

    doc.add_heading("Save costs", level=2)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Shipping total", "4,500"],
            ["Duty total", "1,100"],
            ["VAT total", "1,650"],
            ["Paperwork source", "Agent invoice MW-441 / FR-992"],
        ],
    )

    doc.add_heading("True Cost", level=2)
    add_table(
        doc,
        ["Setting", "Value"],
        [
            ["Allocation", "quantity"],
            ["Profit margin", "0.30 (30%)"],
            ["Action", "Commit batch when ready"],
        ],
    )

    # Shipment 2
    doc.add_heading("Shipment 2 — Auto FX (USD supplier invoice)", level=1)
    doc.add_paragraph(
        "Story: LED lighting from Shenzhen. Supplier invoice is in USD; "
        "freight and tax paperwork are already in ZMW."
    )

    doc.add_heading("Products (create once)", level=2)
    add_table(
        doc,
        ["SKU", "Name", "Weight (kg)"],
        [
            ["LED-9W", "9W LED bulb warm white", "0.08"],
            ["LED-12W", "12W LED bulb daylight", "0.10"],
            ["LED-BATT", "Rechargeable LED lantern", "0.45"],
        ],
    )

    doc.add_heading("New shipment header", level=2)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Reference", "SHP-FX-2026-02"],
            ["Supplier", "BrightPath Electronics Ltd"],
            ["Invoice #", "BP-SZ-11907"],
            ["Clearing agent", "Zed Customs Partners"],
            ["Carrier", "Air China Cargo"],
            ["Origin", "China"],
            ["Primary cost mode", "Auto FX"],
            ["Foreign currency", "USD"],
            ["Local currency", "ZMW"],
        ],
    )

    doc.add_heading("Lines (mode: Auto FX on each line)", level=2)
    add_table(
        doc,
        ["Product", "Qty", "Line weight (kg)", "Foreign unit (USD)"],
        [
            ["LED-9W", "50", "4.0", "0.85"],
            ["LED-12W", "30", "3.0", "1.20"],
            ["LED-BATT", "20", "9.0", "2.50"],
        ],
    )

    doc.add_heading("Save costs", level=2)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Exchange rate", "27.40"],
            ["Shipping total", "6,800"],
            ["Duty total", "2,400"],
            ["VAT total", "3,200"],
            ["Paperwork source", "Supplier inv BP-SZ-11907 / CA slip ZCP-778"],
        ],
    )
    doc.add_paragraph(
        "Note: You can use Fetch FX instead of typing 27.40 if the backend has internet access."
    )

    doc.add_heading("True Cost", level=2)
    add_table(
        doc,
        ["Setting", "Value"],
        [
            ["Allocation", "weight"],
            ["Profit margin", "0.25 (25%)"],
            ["Action", "Commit batch when ready"],
        ],
    )

    doc.add_heading("Quick order of work", level=1)
    doc.add_paragraph("1. Products — add all six SKUs listed above.")
    doc.add_paragraph(
        "2. Shipments — create SHP-LOC-2026-01, add lines, save costs, "
        "Open True Cost, allocate, price, commit."
    )
    doc.add_paragraph("3. Repeat for SHP-FX-2026-02 with USD lines and exchange rate.")
    doc.add_paragraph("")
    p = doc.add_paragraph("Generated for Tranz-Algo local demo and viva practice.")
    p.runs[0].font.size = Pt(10)

    doc.save(OUT)
    print(f"Created: {OUT}")


if __name__ == "__main__":
    main()
