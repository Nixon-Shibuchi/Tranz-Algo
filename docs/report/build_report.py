"""
Build Tranz-Algo Final Year Project Report (Word) following the
University of Zambia 2026 report template structure.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from draw_figures import draw_all

# Print-oriented sizes (A4, Times New Roman, graded hard copy)
BODY_SIZE = 12
CAPTION_SIZE = 11
TABLE_SIZE = 11
H1_SIZE = 16
H2_SIZE = 13
FIGURE_WIDTH = Inches(6.1)

OUT_DIR = Path(__file__).resolve().parent
FIG_DIR = OUT_DIR / "figures"
SHOT_DIR = FIG_DIR / "screenshots"
LOGO_TRANZ = FIG_DIR / "tranzalgo_logo.png"
LOGO_UNZA = FIG_DIR / "unza_logo.png"

# Fill computer numbers / supervisor before final submission if known
SUPERVISOR = "Mrs. Leena Kumar"
DEGREE = "Bachelor of Science in Computer Science"
SUBMISSION_DATE = "October 2026"
# Shared source repository (submitted evidence; each student files their own Word report)
REPO_URL = "https://github.com/Nixon-Shibuchi/Tranz-Algo"
PROJECT_TITLE = (
    "Tranz-Algo: Explainable Landed-Cost Inventory Orchestration "
    "with Dual Cost Modes for Import-Dependent SMEs"
)

AUTHORS = {
    "nixon": {
        "key": "nixon",
        "name": "Nixon Shibuchi",
        "computer": "2021437361",
        "slug": "Nixon_Shibuchi",
        "role_short": "implementation lead",
        "focus": (
            "FastAPI backend, costing engine, most frontend wiring, FX service, "
            "evaluation scenarios/oracles, literature lead drafts, and majority of commits"
        ),
        "results": (
            "Oracle tests and core engine results are primarily attributable to this author’s "
            "implementation work. Design ownership remains joint with the co-author."
        ),
        "appendix_detail": (
            "Primary implementation: FastAPI backend; costing engine and dual-mode formulas; "
            "most frontend integration; FX service; scenarios/oracles; literature lead drafts; "
            "majority of git commits (including pending updates)."
        ),
    },
    "musonda": {
        "key": "musonda",
        "name": "Musonda Katongo",
        "computer": "2020583969",
        "slug": "Musonda_Katongo",
        "role_short": "co-design and UI / True Cost contributor",
        "focus": (
            "joint architecture and technology decisions, literature review, "
            "UI redesign and True Cost workflow implementation, and live demo drive"
        ),
        "results": (
            "UI and True Cost workflow refinements are primarily attributable to this author’s "
            "implementation contribution. Design ownership remains joint with the co-author; "
            "oracle fixtures were co-used for evaluation."
        ),
        "appendix_detail": (
            "Co-design partner for architecture, technology, and formulas before coding; "
            "literature review on related-work drafts; UI redesign and True Cost workflow "
            "implementation commit; viva/demo drive role for the live walkthrough."
        ),
    },
}


def partner_of(author):
    return AUTHORS["musonda"] if author["key"] == "nixon" else AUTHORS["nixon"]


def set_run_font(run, size=BODY_SIZE, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    # Pure black for print (avoid grey / theme colours that wash out)
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def configure_print_styles(doc):
    """Force Normal / Heading styles to print-clear Times New Roman."""
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(8)

    for name, size in (("Heading 1", H1_SIZE), ("Heading 2", H2_SIZE), ("Heading 3", 12)):
        try:
            st = styles[name]
        except KeyError:
            continue
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _clear_paragraph(p):
    for child in list(p._element):
        p._element.remove(child)


def configure_section_geometry(section):
    section.page_width = Cm(21.0)  # A4
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)


def set_page_number_format(section, *, fmt="decimal", start=None):
    """fmt: decimal | lowerRoman | upperRoman."""
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag == qn("w:pgNumType"):
            sectPr.remove(child)
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:fmt"), fmt)
    if start is not None:
        pgNumType.set(qn("w:start"), str(start))
    sectPr.append(pgNumType)


def add_page_number(section, *, prefix=""):
    """Centered PAGE field in the section footer (unlink from previous)."""
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    _clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    if prefix:
        set_run_font(p.add_run(prefix), size=10)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    set_run_font(run, size=10)


def add_para(doc, text, *, size=BODY_SIZE, bold=False, italic=False, align=None, space_after=10, first_line=True):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line and align is None:
        pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    size = H1_SIZE if level == 1 else (H2_SIZE if level == 2 else 12)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(8)
    for run in h.runs:
        set_run_font(run, size=size, bold=True)
    return h


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text)
        set_run_font(r2)
    else:
        set_run_font(p.add_run(text))
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run(text)
    set_run_font(run, size=CAPTION_SIZE, italic=True, bold=True)
    return p


def add_figure(doc, path, caption, width=FIGURE_WIDTH):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run()
    run.add_picture(str(path), width=width)
    add_caption(doc, caption)


def add_table(doc, headers, rows, caption=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, "D9E2E8")  # light slate — prints clearly, not washed out
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                set_run_font(r, size=TABLE_SIZE, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            if ri % 2 == 1:
                set_cell_shading(cell, "F3F6F8")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    set_run_font(r, size=TABLE_SIZE)
    if caption:
        add_caption(doc, caption)
    doc.add_paragraph()


def page_break(doc):
    doc.add_page_break()


def add_toc_entry(doc, text, *, level=1):
    """Printable TOC line (indent by heading level; no invented page numbers)."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.5 * (level - 1))
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(2 if level == 1 else 0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_run_font(p.add_run(text), size=BODY_SIZE, bold=(level == 1))


def build_table_of_contents(doc):
    """Static table of contents matching report headings (visible when printed)."""
    add_heading_styled(doc, "TABLE OF CONTENTS", 1)
    entries = [
        (1, "DECLARATION"),
        (1, "APPROVAL"),
        (1, "ABSTRACT"),
        (1, "ACKNOWLEDGEMENTS"),
        (1, "NOTE ON GROUP PROJECTS"),
        (1, "LIST OF ABBREVIATIONS AND ACRONYMS"),
        (1, "Chapter 1  Introduction"),
        (2, "1.1 Background and Context"),
        (2, "1.2 Problem Statement"),
        (2, "1.3 Project Aim"),
        (2, "1.4 Project Objectives"),
        (2, "1.5 Research Questions"),
        (2, "1.6 Scope of the Project"),
        (2, "1.7 Significance of the Project"),
        (2, "1.8 Limitations"),
        (2, "1.9 Organisation of the Report"),
        (1, "Chapter 2  Literature Review"),
        (2, "2.1 Introduction"),
        (2, "2.2 Key Concepts and Theories"),
        (2, "2.3 Existing Systems / Related Work"),
        (2, "2.4 Comparative Analysis"),
        (2, "2.5 Research / Knowledge Gap"),
        (2, "2.6 Conceptual Framework"),
        (2, "2.7 Chapter Summary"),
        (1, "Chapter 3  Methodology"),
        (2, "3.1 Introduction"),
        (2, "3.2 Research / Development Approach"),
        (2, "3.3 Requirements / Data Collection"),
        (2, "3.4 System Architecture / Research Workflow"),
        (2, "3.5 Tools and Technologies"),
        (2, "3.6 Ethical and Legal Considerations"),
        (2, "3.7 Evaluation Method"),
        (2, "3.8 Chapter Summary"),
        (1, "Chapter 4  System Analysis and Design"),
        (2, "4.1 Introduction"),
        (2, "4.2 Stakeholder and User Analysis"),
        (2, "4.3 Functional Requirements"),
        (2, "4.4 Non-Functional Requirements"),
        (2, "4.5 Use Cases / User Stories"),
        (2, "4.6 System / Solution Architecture"),
        (2, "4.7 Database / Data Design"),
        (2, "4.8 Interface / Interaction Design"),
        (2, "4.9 Algorithms / Models / Technical Design"),
        (2, "4.10 Chapter Summary"),
        (1, "Chapter 5  Implementation / Development"),
        (2, "5.1 Introduction"),
        (2, "5.2 Development Environment"),
        (2, "5.3 System Components / Modules"),
        (2, "5.4 Key Implementation Details"),
        (2, "5.5 Integration"),
        (2, "5.6 Challenges and Solutions"),
        (2, "5.7 Individual Contribution (group projects)"),
        (2, "5.8 Chapter Summary"),
        (1, "Chapter 6  Testing, Results and Evaluation"),
        (2, "6.1 Introduction"),
        (2, "6.2 Testing Strategy"),
        (2, "6.3 Test Cases and Results"),
        (2, "6.4 Experimental Results"),
        (2, "6.5 User Evaluation"),
        (2, "6.6 Analysis of Results"),
        (2, "6.7 Objective-by-Objective Evaluation"),
        (2, "6.8 Individual Results"),
        (2, "6.9 Chapter Summary"),
        (1, "Chapter 7  Discussion"),
        (2, "7.1 Introduction"),
        (2, "7.2 Interpretation of Findings"),
        (2, "7.3 Comparison with Existing Work"),
        (2, "7.4 Contributions of the Project"),
        (2, "7.5 Limitations and Implications"),
        (2, "7.6 Chapter Summary"),
        (1, "Chapter 8  Conclusion and Recommendations"),
        (2, "8.1 Introduction"),
        (2, "8.2 Summary of the Project"),
        (2, "8.3 Conclusions"),
        (2, "8.4 Achievement of Objectives"),
        (2, "8.5 Recommendations"),
        (2, "8.6 Future Work"),
        (1, "REFERENCES"),
        (1, "Appendix A  Individual Contribution Statement"),
        (2, "A.1 Project and Group Information"),
        (2, "A.2 Overall Group Project"),
        (2, "A.3 Individual Role and Responsibilities"),
        (2, "A.4 Collaboration and Co-author Work"),
        (1, "Appendix B  User Manual / Installation Guide (summary)"),
    ]
    for level, text in entries:
        add_toc_entry(doc, text, level=level)
    page_break(doc)


def add_centered_picture(doc, path, width=Inches(1.45)):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run()
    run.add_picture(str(path), width=width)


def build_cover(doc, author):
    partner = partner_of(author)
    # Logos: UNZA + Tranz-Algo
    logos = doc.add_paragraph()
    logos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logos.paragraph_format.first_line_indent = Cm(0)
    logos.paragraph_format.space_after = Pt(10)
    if LOGO_UNZA.exists():
        logos.add_run().add_picture(str(LOGO_UNZA), width=Inches(1.35))
    logos.add_run("    ")
    if LOGO_TRANZ.exists():
        logos.add_run().add_picture(str(LOGO_TRANZ), width=Inches(1.35))

    for line, size, bold in [
        ("THE UNIVERSITY OF ZAMBIA", 16, True),
        ("SCHOOL OF NATURAL AND APPLIED SCIENCES", 13, True),
        ("DEPARTMENT OF COMPUTING AND INFORMATICS", 13, True),
        ("", 12, False),
        ("FINAL YEAR PROJECT REPORT", 16, True),
        ("ACADEMIC YEAR 2026", 13, True),
        ("", 12, False),
        (PROJECT_TITLE, 14, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        if line:
            set_run_font(p.add_run(line), size=size, bold=bold)
    doc.add_paragraph()
    for line in [
        "A Final Year Project Report Submitted to the",
        "Department of Computing and Informatics",
        "in partial fulfilment of the requirements for the award of",
        DEGREE,
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        set_run_font(p.add_run(line), size=12)

    doc.add_paragraph()
    meta = [
        f"Student Name: {author['name']}",
        f"Computer Number: {author['computer']}",
        f"Group project with: {partner['name']} ({partner['computer']})",
        f"Supervisor: {SUPERVISOR}",
        f"Date of Submission: {SUBMISSION_DATE}",
        f"Source code repository: {REPO_URL}",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        set_run_font(p.add_run(line), size=12)
    page_break(doc)


def build_front_matter(doc, author):
    partner = partner_of(author)
    add_heading_styled(doc, "DECLARATION", 1)
    add_para(
        doc,
        "I declare that this final year project report is my own work and has been prepared in "
        "accordance with the academic requirements of the University of Zambia. All sources of "
        "information, ideas, data, software, images, and other material that are not my own have been "
        "appropriately acknowledged and referenced.",
        first_line=False,
    )
    add_para(
        doc,
        "Where this project was undertaken as a group, I have clearly identified my individual "
        "contributions and have not presented the work of other group members as my own. "
        f"Tranz-Algo was developed with {partner['name']}. Design and technology decisions were "
        f"made jointly. My primary role was {author['role_short']}: {author['focus']}. "
        "See Appendix A.",
        first_line=False,
    )
    add_para(doc, f"Student Name: {author['name']}", first_line=False)
    add_para(doc, f"Computer Number: {author['computer']}", first_line=False)
    add_para(doc, "Signature: ______________________", first_line=False)
    add_para(doc, f"Date: {SUBMISSION_DATE}", first_line=False)
    page_break(doc)

    add_heading_styled(doc, "APPROVAL", 1)
    add_para(
        doc,
        "This final year project report has been submitted for examination with the approval of the "
        "project supervisor.",
        first_line=False,
    )
    add_para(doc, f"Supervisor Name: {SUPERVISOR}", first_line=False)
    add_para(doc, "Signature: ______________________", first_line=False)
    add_para(doc, "Date: ______________________", first_line=False)
    add_para(doc, "Departmental/Examiner Approval (if applicable): ______________________", first_line=False)
    page_break(doc)

    add_heading_styled(doc, "ABSTRACT", 1)
    add_para(
        doc,
        "Import-dependent retailers often know supplier invoices and clearing documents yet still "
        "cannot state a defensible per-item true cost. Shared freight and tax, mixed foreign and local "
        "paperwork, and exchange-rate movement undermine cost-plus pricing. This project presents "
        "Tranz-Algo, a web-based decision-support system that computes explainable landed costs for "
        "import shipment batches, suggests retail prices, and records an auditable Cost history.",
        first_line=False,
    )
    add_para(
        doc,
        "Tranz-Algo implements dual cost-entry modes (Auto FX and Already-Local), quantity and weight "
        "allocation with a comparator, a True Cost pricing studio with a safety floor, batch-linked "
        "stock, and a lightweight Record sale that reduces stock without becoming a POS. The system "
        "uses FastAPI, a React (Vite) administrative UI, PostgreSQL with SQLite fallback, and free "
        "Frankfurter exchange rates with cache and manual override. Evaluation against spreadsheet "
        "oracles for two demo shipments (Already-Local grocery; Auto FX phone parts) passed within "
        "±0.01 ZMW. The contribution is explainable import-batch costing under mixed paperwork, not "
        "another inventory CRUD or full ERP.",
        first_line=False,
    )
    add_para(
        doc,
        f"This individual report is prepared by {author['name']} for the group project developed "
        f"with {partner['name']}. Only this individual report is submitted for examination; "
        f"the shared implementation is available at {REPO_URL}.",
        first_line=False,
    )
    add_para(
        doc,
        "Keywords: landed cost; dual cost modes; import SME; allocation; True-Cost Ledger; FastAPI; React",
        first_line=False,
        italic=True,
    )
    page_break(doc)

    add_heading_styled(doc, "ACKNOWLEDGEMENTS", 1)
    add_para(
        doc,
        f"I thank my supervisor, {SUPERVISOR}, for guidance throughout this final year project. "
        "I also thank the Department of Computing and Informatics, University of Zambia, for "
        "providing the academic framework and assessment standards that shaped this work. Finally, "
        f"I acknowledge {partner['name']} for joint design discussions and complementary "
        "implementation effort that produced the working Tranz-Algo MVP.",
        first_line=False,
    )
    page_break(doc)

    add_heading_styled(doc, "NOTE ON GROUP PROJECTS", 1)
    add_para(
        doc,
        "This is an individual report for a two-member group project. The shared system is Tranz-Algo. "
        f"The co-author is {partner['name']}. Each student submits only their own individual report "
        "(not a joint Word document). The shared source code and documentation live in the project "
        f"Git repository: {REPO_URL}. Individual responsibilities for {author['name']} are stated "
        "in Chapter 5 and Appendix A. Each student remains responsible for understanding the full "
        "system and for the honesty of their contribution claims.",
        first_line=False,
    )
    page_break(doc)


def chapter_1(doc):
    add_heading_styled(doc, "Chapter 1  Introduction", 1)
    add_heading_styled(doc, "1.1 Background and Context", 2)
    add_para(
        doc,
        "Zambian and other import-dependent SMEs buy goods in shipment batches. Supplier invoices may "
        "be in foreign currency, while freight, duty, and VAT arrive as local clearing totals. Retailers "
        "must recover these shared costs per SKU before setting a selling price. Spreadsheets are common "
        "but fragile; POS tools sell well but rarely explain import true cost; ERP landed-cost modules "
        "are powerful but heavy for small shops.",
    )
    add_para(
        doc,
        "Tranz-Algo targets the decision-support niche between clearing paperwork and the retail "
        "catalogue: compute landed cost transparently, suggest a safe retail price, commit an audit "
        "trail, and update stock by batch.",
    )

    add_heading_styled(doc, "1.2 Problem Statement", 2)
    add_para(
        doc,
        "Import retailers frequently cannot defend a per-item true cost when freight and tax are shared, "
        "paperwork mixes foreign and local amounts, and exchange rates move after prices are set. "
        "Existing tools either omit explainable allocation, hide formulas in spreadsheets, or require "
        "full ERP complexity. The problem addressed is therefore explainable import-batch costing and "
        "price suggestion under mixed real-world paperwork.",
    )

    add_heading_styled(doc, "1.3 Project Aim", 2)
    add_para(
        doc,
        "To design, implement, and evaluate a free-to-run web system that turns import shipment "
        "paperwork into per-item landed cost, suggested retail price, auditable Cost history, and "
        "batch-linked inventory for SME administrators.",
    )

    add_heading_styled(doc, "1.4 Project Objectives", 2)
    for o in [
        "Analyse import costing needs and freeze dual-mode formulas (Auto FX and Already-Local).",
        "Design a three-tier architecture, data model, and True Cost workflow with allocation comparison.",
        "Implement FastAPI services, React admin UI, PostgreSQL persistence, and FX integration.",
        "Evaluate correctness against spreadsheet oracles (tolerance ≤ 0.01 ZMW) using demo scenarios.",
        "Document technology deviations from the proposal and produce operator / demo guidance.",
    ]:
        add_bullet(doc, o)

    add_heading_styled(doc, "1.5 Research Questions", 2)
    add_bullet(doc, "How can dual entry modes accept mixed foreign/local import paperwork without forcing one workflow?")
    add_bullet(doc, "How can allocation of shared shipping/duty/VAT be made explainable to an SME user?")
    add_bullet(doc, "Can automated landed-cost and price suggestion match a spreadsheet oracle within ±0.01 ZMW?")

    add_heading_styled(doc, "1.6 Scope of the Project", 2)
    add_heading_styled(doc, "1.6.1 In Scope", 3)
    add_bullet(doc, "Admin authentication (JWT); product catalogue; shipment batches and lines.")
    add_bullet(doc, "Dual modes, allocation (quantity/weight), True Cost studio, safety floor, Cost history.")
    add_bullet(doc, "CSV import for demos; stock board; Record sale (stock-out only).")
    add_bullet(doc, "Frankfurter FX with cache/manual; PostgreSQL primary and SQLite fallback.")

    add_heading_styled(doc, "1.6.2 Out of Scope", 3)
    add_bullet(doc, "Full POS checkout (payments, customers, receipts).")
    add_bullet(doc, "Legal ZRA/customs duty assessment APIs.")
    add_bullet(doc, "React Native companion app (deferred; web admin delivered).")
    add_bullet(doc, "Multi-tenant SaaS billing; barcode camera as a core path.")

    add_heading_styled(doc, "1.7 Significance of the Project", 2)
    add_para(
        doc,
        "Practically, Tranz-Algo helps import SMEs defend prices with explainable costs. Academically, "
        "it demonstrates that the gap is not inventory CRUD but explainability under mixed paperwork. "
        "Institutionally, it provides a complete design–build–evaluate cycle with reproducible demos.",
    )

    add_heading_styled(doc, "1.8 Limitations", 2)
    add_para(
        doc,
        "Auto FX seeds the rate at **20** (or the last fetched cache) until the user clicks Fetch FX "
        "or types a rate. Demo B spreadsheet/pytest oracles use the same **20** so the UI starter and "
        "evaluation rate match. Live Fetch FX may differ; keep 20 for the frozen demo. Value-based "
        "allocation exists in the engine but is not primary in the UI. The MVP assumes honest paperwork "
        "totals and does not replace clearing agents.",
    )

    add_heading_styled(doc, "1.9 Organisation of the Report", 2)
    add_para(
        doc,
        "Chapter 2 reviews literature. Chapter 3 describes methodology. Chapter 4 presents analysis "
        "and design. Chapter 5 covers implementation. Chapter 6 reports testing and evaluation. "
        "Chapter 7 discusses findings. Chapter 8 concludes. Appendices include the contribution "
        "statement and installation notes.",
        first_line=False,
    )


def chapter_2(doc):
    add_heading_styled(doc, "Chapter 2  Literature Review", 1)
    add_heading_styled(doc, "2.1 Introduction", 2)
    add_para(
        doc,
        "This chapter reviews landed-cost concepts, FX risk for SMEs, competing tool classes, and the "
        "knowledge gap that Tranz-Algo addresses. Sources are listed in the References.",
    )

    add_heading_styled(doc, "2.2 Key Concepts and Theories", 2)
    add_para(
        doc,
        "Landed cost is broader than supplier invoice price: product, freight, customs, and related "
        "charges jointly determine unit economics (Mangan & Lalwani, 2016; Bonanni, 2011; NetSuite). "
        "Allocation bases commonly include quantity, weight, or value. Cost-plus pricing remains a "
        "practical SME posture: Suggested RP = Landed × (1 + margin) (Nagle & Müller). Exchange-rate "
        "volatility can erase margins for import-dependent firms that do not hedge (DiVA study; "
        "IJRISS, 2025 on Zambian MSMEs).",
    )

    add_heading_styled(doc, "2.3 Existing Systems / Related Work", 2)
    add_para(
        doc,
        "POS applications optimise selling but rarely allocate shared import charges with an audit "
        "trail. Basic inventory tools track counts more than explainable landed cost. Spreadsheets can "
        "encode any formula yet are opaque and fragile. Shipping/clearing systems produce paperwork "
        "but do not price a retail catalogue. ERP suites (for example Odoo landed costs) can implement "
        "sophisticated costing but are heavy for many small importers.",
    )

    add_heading_styled(doc, "2.4 Comparative Analysis", 2)
    add_table(
        doc,
        ["System type", "Import batch true cost?", "Dual modes?", "Explainable allocation?", "Fit for small importer"],
        [
            ["POS / shop apps", "Rarely", "No", "No", "Selling-focused"],
            ["Basic inventory", "Weak / manual", "Usually no", "Usually no", "Quantity-focused"],
            ["Spreadsheets", "Possible, fragile", "Manual", "Opaque", "Common but error-prone"],
            ["Clearing / shipping", "Paperwork only", "Partial", "No retail pricing", "Logistics partner"],
            ["ERP (Odoo/SAP-class)", "Often yes", "Varies", "Possible, heavy", "Often too complex"],
            ["Tranz-Algo", "Yes (core)", "Yes", "Yes (comparator + ledger)", "Built for this niche"],
        ],
        "Table 2.1: Comparison of tool classes against Tranz-Algo’s niche",
    )

    add_heading_styled(doc, "2.5 Research / Knowledge Gap", 2)
    add_para(
        doc,
        "The gap is not another inventory app. It is explainable import-batch costing and price "
        "suggestion under mixed real-world paperwork. Dual modes accept foreign invoices and "
        "already-local shipper amounts. An Allocation Comparator and Cost history / True-Cost Ledger "
        "make methods visible and auditable.",
    )

    add_heading_styled(doc, "2.6 Conceptual Framework", 2)
    add_para(
        doc,
        "Tranz-Algo follows a three-tier decision-support framing: capture paperwork → apply frozen "
        "costing strategies → persist explainable outcomes to stock and ledger. Engineering guidance "
        "emphasises verifying critical financial logic (Pressman & Maxim, 2020).",
    )
    add_figure(doc, FIG_DIR / "dual_modes.png", "Figure 2.1: Dual cost-entry modes conceptual framework")

    add_heading_styled(doc, "2.7 Chapter Summary", 2)
    add_para(
        doc,
        "Literature supports multi-component landed cost, FX-aware SME pricing, and a niche between "
        "paperwork and POS/ERP. Tranz-Algo is positioned to fill that niche with dual modes and auditability.",
    )


def chapter_3(doc):
    add_heading_styled(doc, "Chapter 3  Methodology", 1)
    add_heading_styled(doc, "3.1 Introduction", 2)
    add_para(
        doc,
        "The project combined requirements freeze for costing formulas with incremental vertical "
        "delivery of working software, followed by oracle-based evaluation.",
    )

    add_heading_styled(doc, "3.2 Research / Development Approach", 2)
    add_para(
        doc,
        "The proposal named Waterfall for early clarity on financial formulas. In practice, formulas "
        "were frozen first, then modules were delivered in slices: auth → products → shipments → "
        "True Cost → stock. This preserved correctness discipline while allowing feedback on UX.",
    )

    add_heading_styled(doc, "3.3 Requirements / Data Collection", 2)
    add_para(
        doc,
        "Requirements came from the approved proposal, joint design discussions, and review of "
        "import paperwork patterns (foreign invoices vs local clearing totals). Evaluation data are "
        "synthetic but realistic demo shipments (grocery Already-Local; phone-parts Auto FX) with "
        "spreadsheet oracles in docs/scenarios/.",
    )

    add_heading_styled(doc, "3.4 System Architecture / Research Workflow", 2)
    add_para(
        doc,
        "Architecture is three-tier: React frontend → FastAPI /api/v1 → PostgreSQL (SQLite fallback), with "
        "optional FX providers. The True Cost workflow is allocate → suggest/adjust → commit → Cost history.",
    )
    add_figure(doc, FIG_DIR / "architecture.png", "Figure 3.1: System architecture used in the methodology")

    add_heading_styled(doc, "3.5 Tools and Technologies", 2)
    add_table(
        doc,
        ["Area", "Choice"],
        [
            ["Languages", "Python 3, JavaScript (React)"],
            ["Backend", "FastAPI, SQLAlchemy 2, Pydantic, JWT/bcrypt"],
            ["Frontend", "React (Vite), CSS themes"],
            ["Database", "PostgreSQL primary; SQLite fallback"],
            ["FX", "Frankfurter / open.er-api + cache/manual"],
            ["Tests", "pytest costing oracles"],
            ["VCS", "Git / GitHub"],
            ["OS (dev)", "Windows 10/11"],
        ],
        "Table 3.1: Tools and technologies",
    )

    add_heading_styled(doc, "3.6 Ethical and Legal Considerations", 2)
    add_para(
        doc,
        "User passwords are hashed (bcrypt). The system does not claim legal customs assessment. Demo "
        "data are synthetic. Free FX APIs are used under their published terms. Group contribution is "
        "declared honestly.",
    )

    add_heading_styled(doc, "3.7 Evaluation Method", 2)
    add_para(
        doc,
        "Success is measured by (1) formula agreement with spreadsheet oracles within ±0.01 ZMW, "
        "(2) scenario coverage of dual modes, allocation comparison, safety floor, and UI CSV path, "
        "and (3) demonstration that Record sale is stock-out only (not POS).",
    )

    add_heading_styled(doc, "3.8 Chapter Summary", 2)
    add_para(
        doc,
        "Methodology froze costing rules, implemented a three-tier MVP, and evaluated against reproducible oracles.",
    )


def chapter_4(doc):
    add_heading_styled(doc, "Chapter 4  System Analysis and Design", 1)
    add_heading_styled(doc, "4.1 Introduction", 2)
    add_para(doc, "This chapter presents stakeholders, requirements, architecture, data design, interfaces, and algorithms.")

    add_heading_styled(doc, "4.2 Stakeholder and User Analysis", 2)
    add_bullet(doc, "Primary user: import SME administrator (desk workflow).")
    add_bullet(doc, "Indirect: shop staff who rely on defended retail prices and stock levels.")
    add_bullet(doc, "Out of scope users: end customers at till; customs officers.")

    add_heading_styled(doc, "4.3 Functional Requirements", 2)
    for fr in [
        "FR1 Register/sign in with JWT.",
        "FR2 Manage products (including CSV import and optional auto SKU).",
        "FR3 Create shipments in Already-Local or Auto FX mode; enter lines and paperwork totals.",
        "FR4 Fetch/cache/manual FX for Auto FX.",
        "FR5 Allocate by quantity or weight; compare methods.",
        "FR6 Suggest retail prices; apply safety floor while adjusting; commit with Cost history.",
        "FR7 Update stock on commit; Record sale reduces stock only.",
        "FR8 Admin Start fresh / Clear drafts for demos.",
    ]:
        add_bullet(doc, fr)

    add_heading_styled(doc, "4.4 Non-Functional Requirements", 2)
    add_bullet(doc, "NFR1 Correctness: money calculations within ±0.01 ZMW of oracle.")
    add_bullet(doc, "NFR2 Security: bcrypt hashes; JWT on protected routes.")
    add_bullet(doc, "NFR3 Usability: clear True Cost stages; Create account (not create admin).")
    add_bullet(doc, "NFR4 Deployability: PostgreSQL via DATABASE_URL; SQLite fallback.")
    add_bullet(doc, "NFR5 Maintainability: Strategy/Facade patterns for modes and commit.")

    add_heading_styled(doc, "4.5 Use Cases / User Stories", 2)
    add_bullet(doc, "As an admin, I import Demo A CSV and cost an Already-Local grocery batch.")
    add_bullet(doc, "As an admin, I open Auto FX at rate 20 (CSV / UI starter), allocate by weight, and cost Demo B.")
    add_bullet(doc, "As an admin, I review Cost history after commit and record a sale that reduces stock.")

    add_heading_styled(doc, "4.6 System / Solution Architecture", 2)
    add_para(
        doc,
        "Figure 4.1 shows the as-built architecture. Proposal deltas: React web instead of React Native; "
        "PostgreSQL local/Docker instead of mandatory Supabase; on-demand FX instead of continuous streaming.",
    )
    add_figure(doc, FIG_DIR / "architecture.png", "Figure 4.1: Tranz-Algo system architecture (as built)")

    add_heading_styled(doc, "4.7 Database / Data Design", 2)
    add_para(
        doc,
        "Core entities: users, products, shipment_batches, batch_items, cost_entries, ledger_entries, "
        "sales, fx_rate_cache. Tax profiles were removed; duty/VAT live as batch paperwork totals. "
        "Primary keys are integer serials.",
    )
    add_figure(doc, FIG_DIR / "erd.png", "Figure 4.2: UML-style class / entity diagram (MVP)")

    add_heading_styled(doc, "4.8 Interface / Interaction Design", 2)
    add_para(
        doc,
        "The UI is a React frontend with Welcome, Login/Create account, Dashboard, Products, Shipments "
        "(CSV import, Clear drafts), Shipment detail (Fetch FX), True Cost studio (Lines → Allocate → "
        "Price → Commit), Stock, and Cost history (When / What). Themes and a closable sidebar support "
        "desk use. Figures 4.4–4.8 are as-built screenshots of the working MVP.",
    )
    add_figure(doc, FIG_DIR / "true_cost_workflow.png", "Figure 4.3: True Cost interaction flow")

    # As-built UI screenshots (from docs/Screenshots/)
    ui_shots = [
        (
            "ui_dashboard.png",
            "Figure 4.4: Dashboard — stock and shipment overview with Start fresh and Open True Cost",
            "The Dashboard summarises units available, low/out counts, draft versus committed batches, "
            "and a short on-the-shelf list so the administrator can see status before opening True Cost.",
        ),
        (
            "ui_products.png",
            "Figure 4.5: Products — catalogue entry form and CSV import",
            "Products supports manual create (optional auto SKU) and Import CSV for demo catalogues. "
            "Optional details are folded away to keep the primary form short.",
        ),
        (
            "ui_shipments.png",
            "Figure 4.6: Shipments — create batch with dual cost mode (Already-Local / Auto FX)",
            "Primary cost mode is chosen once when the shipment is created; later lines and CSV imports "
            "follow that mode. Clear drafts and Import shipment CSV sit in the page actions.",
        ),
        (
            "ui_true_cost_pricing.png",
            "Figure 4.7: True Cost — Step 2 retail suggestion with margin slider and safety floor",
            "After allocation, the administrator drags profit margin; suggested retail prices update live. "
            "The red mark is the safety floor (10% above landed unit cost).",
        ),
        (
            "ui_stock.png",
            "Figure 4.8: Stock — available / low / out board with Record sale",
            "Stock shows units on hand by SKU with retail context and a Record sale action that reduces "
            "quantity without becoming a full POS.",
        ),
    ]
    for filename, caption, blurb in ui_shots:
        path = SHOT_DIR / filename
        if not path.exists():
            continue
        add_para(doc, blurb)
        add_figure(doc, path, caption)

    add_heading_styled(doc, "4.9 Algorithms / Models / Technical Design", 2)
    add_para(doc, "Auto FX product local = foreign_unit_price × exchange_rate.", first_line=False)
    add_para(doc, "Already-Local product local = local_unit_product_cost.", first_line=False)
    add_para(
        doc,
        "Shared shipping/duty/VAT shares use quantity or weight proportions. "
        "Landed_unit = (product_local_line + shares) / qty. "
        "Suggested_RP = Landed_unit × (1 + margin). "
        "Safety floor = Landed_unit × (1.10) on live True Cost adjust (no post-commit recheck UI).",
        first_line=False,
    )
    add_para(
        doc,
        "Design patterns: Strategy for cost modes/allocation; Facade for allocate/suggest/commit; "
        "Chain of Responsibility for FX providers.",
    )

    add_heading_styled(doc, "4.10 Chapter Summary", 2)
    add_para(doc, "Analysis and design freeze dual modes, three-tier architecture, ERD, and True Cost algorithms for implementation.")


def chapter_5(doc, author):
    partner = partner_of(author)
    add_heading_styled(doc, "Chapter 5  Implementation / Development", 1)
    add_heading_styled(doc, "5.1 Introduction", 2)
    add_para(doc, "This chapter describes the working MVP, environment, modules, key decisions, integration, challenges, and individual contributions.")

    add_heading_styled(doc, "5.2 Development Environment", 2)
    add_para(
        doc,
        "Development used Windows with Python virtualenv, Node/Vite for the frontend, PostgreSQL "
        "(Docker or local install) via DATABASE_URL, and Git. Backend runs on port 8000; frontend on 5173.",
    )

    add_heading_styled(doc, "5.3 System Components / Modules", 2)
    add_table(
        doc,
        ["Module", "Backend", "Frontend"],
        [
            ["Auth", "api/v1/auth.py", "LoginPage / Create account"],
            ["Products", "api/v1/products.py", "ProductsPage + CSV"],
            ["Shipments", "api/v1/shipments.py", "Shipments / Detail + CSV"],
            ["True Cost", "shipment_workflow.py, costing.py", "TrueCostPage"],
            ["FX", "fx.py, fx_providers.py", "Fetch FX on costs"],
            ["Stock", "inventory.py", "InventoryPage / Record sale"],
            ["Admin", "admin.py", "Dashboard Start fresh"],
        ],
        "Table 5.1: Implemented modules",
    )

    add_heading_styled(doc, "5.4 Key Implementation Details", 2)
    add_para(
        doc,
        "PRIMARY_COST_MODES exposes Already-Local and Auto FX only. Duty/VAT are paperwork totals on "
        "the batch. Auto FX seeds rate from last cache or starter 20 until Fetch FX/manual. Commit "
        "writes ledger payloads for Cost history and increments stock. Record sale decrements stock "
        "without payments or receipts.",
    )

    add_heading_styled(doc, "5.5 Integration", 2)
    add_para(
        doc,
        "The React frontend calls JSON APIs with Bearer JWT. SQLAlchemy maps models to PostgreSQL/SQLite. FX "
        "providers are chained with fallback. CSV import parses demo samples into products and shipments.",
    )

    add_heading_styled(doc, "5.6 Challenges and Solutions", 2)
    add_bullet(doc, "Challenge: mixed paperwork. Solution: dual modes chosen at shipment create.")
    add_bullet(doc, "Challenge: React Native scope risk. Solution: deliver React web admin; defer RN.")
    add_bullet(doc, "Challenge: FX API downtime. Solution: cache, starter rate, manual override.")
    add_bullet(doc, "Challenge: opaque spreadsheets. Solution: comparator + Cost history When/What.")

    add_heading_styled(doc, "5.7 Individual Contribution (group projects)", 2)
    add_para(
        doc,
        "Almost all architecture, technology, and formula decisions were discussed jointly before "
        f"coding. This author’s role ({author['name']}) was {author['role_short']}: {author['focus']}. "
        f"The co-author ({partner['name']}) contributed complementary work as described in Appendix A. "
        "Neither member claims the other’s exclusive deliverables as sole personal work.",
    )

    add_heading_styled(doc, "5.8 Chapter Summary", 2)
    add_para(doc, "The MVP realises the designed dual-mode True Cost system on FastAPI + React + PostgreSQL.")


def chapter_6(doc, author):
    add_heading_styled(doc, "Chapter 6  Testing, Results and Evaluation", 1)
    add_heading_styled(doc, "6.1 Introduction", 2)
    add_para(doc, "Evaluation verifies that Tranz-Algo’s costing matches spreadsheet oracles and that demo UI paths work.")

    add_heading_styled(doc, "6.2 Testing Strategy", 2)
    add_bullet(doc, "Unit/engine tests: pytest fixtures for Demo A/B oracles.")
    add_bullet(doc, "Scenario scripts 01–05: modes, comparator, safety floor, UI CSV path.")
    add_bullet(doc, "Manual system demo: Start fresh → import CSV → True Cost → commit → Cost history → Record sale.")

    add_heading_styled(doc, "6.3 Test Cases and Results", 2)
    add_table(
        doc,
        ["ID", "Test case", "Expected", "Status"],
        [
            ["TC-01", "Already-Local Demo A qty oracles", "Landed/RP within ±0.01", "PASS"],
            ["TC-02", "Auto FX Demo B (rate 20)", "Landed/RP within ±0.01", "PASS"],
            ["TC-03", "Allocation compare Demo A", "Qty vs weight deltas match arithmetic", "PASS"],
            ["TC-04", "Safety floor on True Cost", "RICE floor 187; warn if below", "PASS"],
            ["TC-05", "CSV UI path (Scenario 05)", "Same batches via import", "PASS"],
            ["TC-06", "Record sale stock-out", "Stock decreases; no POS artefacts", "PASS"],
        ],
        "Table 6.1: Representative test cases and results",
    )

    add_heading_styled(doc, "6.4 Experimental Results", 2)
    add_para(
        doc,
        "Demo A (SHP-DEMO-LOCAL-ZMW) locks grocery landed costs such as 170 / 107 / 160 / 120 / 93 with "
        "30% suggested retail. Demo B (SHP-DEMO-AUTO-FX) uses FX rate **20** (UI starter = oracle) and "
        "locks weight-based landed costs 136.25 / 145 / 46.25 / 92.50 (earbud suggested retail 170.31). "
        "CI fixtures in backend/tests/test_costing.py encode these oracles.",
    )

    add_heading_styled(doc, "6.5 User Evaluation", 2)
    add_para(
        doc,
        "Formal large-scale user surveys were out of scope. Evaluation focused on reproducible "
        "correctness and a classroom demo path documented under docs/scenarios/.",
    )

    add_heading_styled(doc, "6.6 Analysis of Results", 2)
    add_para(
        doc,
        "Passing oracles shows the costing engine is faithful to the frozen formulas. Dual-mode "
        "scenarios show both paperwork styles are supported. The safety-floor test confirms "
        "pre-commit margin protection without claiming automated post-commit recheck.",
    )

    add_heading_styled(doc, "6.7 Objective-by-Objective Evaluation", 2)
    add_bullet(doc, "Formulas/analysis: reflected in design docs and tests — achieved.")
    add_bullet(doc, "Architecture/design: Figures 4.1–4.3 — achieved.")
    add_bullet(doc, "Implementation: running MVP — achieved.")
    add_bullet(doc, "Evaluation ≤0.01 ZMW — achieved for demos.")
    add_bullet(doc, "Operator and demo documentation — achieved.")

    add_heading_styled(doc, "6.8 Individual Results", 2)
    add_para(doc, author["results"])

    add_heading_styled(doc, "6.9 Chapter Summary", 2)
    add_para(doc, "Against demo spreadsheet baselines, core true-cost and suggestion calculations are accurate within tolerance.")


def chapter_7(doc):
    add_heading_styled(doc, "Chapter 7  Discussion", 1)
    add_heading_styled(doc, "7.1 Introduction", 2)
    add_para(doc, "This chapter interprets results relative to the problem, literature, and contributions.")

    add_heading_styled(doc, "7.2 Interpretation of Findings", 2)
    add_para(
        doc,
        "Results indicate that a narrow, explainable costing system can match spreadsheet accuracy "
        "while improving auditability. Dual modes address the real paperwork split that single-workflow "
        "tools ignore.",
    )

    add_heading_styled(doc, "7.3 Comparison with Existing Work", 2)
    add_para(
        doc,
        "Unlike POS or basic inventory apps, Tranz-Algo centres import batch true cost. Unlike ERP, "
        "it stays SME-scoped. Unlike spreadsheets, it freezes formulas and stores Cost history.",
    )

    add_heading_styled(doc, "7.4 Contributions of the Project", 2)
    add_bullet(doc, "Working dual-mode True Cost MVP with comparator and ledger.")
    add_bullet(doc, "Justified technology refinements (web admin; Postgres + SQLite fallback; on-demand FX).")
    add_bullet(doc, "Reproducible evaluation artefacts and operator documentation.")

    add_heading_styled(doc, "7.5 Limitations and Implications", 2)
    add_para(
        doc,
        "The system trusts user-entered paperwork; it does not assess duty legally. Live FX must be "
        "pinned for reproducible demos. Broader SME field trials would strengthen usability claims.",
    )

    add_heading_styled(doc, "7.6 Chapter Summary", 2)
    add_para(doc, "Findings support the niche claim: explainable import costing under mixed paperwork.")


def chapter_8(doc):
    add_heading_styled(doc, "Chapter 8  Conclusion and Recommendations", 1)
    add_heading_styled(doc, "8.1 Introduction", 2)
    add_para(doc, "This chapter summarises the project, conclusions, objective achievement, recommendations, and future work.")

    add_heading_styled(doc, "8.2 Summary of the Project", 2)
    add_para(
        doc,
        "Tranz-Algo was designed and built to convert import paperwork into per-item landed cost, "
        "suggested retail, Cost history, and stock. Dual modes, allocation comparison, and a safety "
        "floor address SME pricing under mixed documents and FX pressure.",
    )

    add_heading_styled(doc, "8.3 Conclusions", 2)
    add_para(
        doc,
        "A focused three-tier web system can deliver explainable landed-cost decision support without "
        "becoming POS or ERP. Oracle evaluation confirms formula fidelity for the demo shipments.",
    )

    add_heading_styled(doc, "8.4 Achievement of Objectives", 2)
    add_para(doc, "All stated objectives in Section 1.4 were achieved for the MVP scope.", first_line=False)

    add_heading_styled(doc, "8.5 Recommendations", 2)
    add_bullet(doc, "Keep demo rates pinned when presenting Auto FX publicly.")
    add_bullet(doc, "Train users that duty/VAT are paperwork totals, not legal assessments.")
    add_bullet(doc, "Use Cost history in disputes about how a retail price was derived.")

    add_heading_styled(doc, "8.6 Future Work", 2)
    add_bullet(doc, "Optional React Native companion; cloud-hosted Postgres for off-campus demos.")
    add_bullet(doc, "Richer value-based allocation UI; scheduled margin alerts.")
    add_bullet(doc, "One-click demo seed and structured SME usability study.")


def references(doc):
    add_heading_styled(doc, "REFERENCES", 1)
    refs = [
        "Bonanni, A. (2011). Total Landed Cost for Tactical Supply Chain Decisions. M.Sc. thesis, Politecnico di Milano.",
        "Frankfurter API documentation (ECB reference rates). https://www.frankfurter.app/docs/",
        "IJERT (2023). Maximizing Profitability through Landed Cost Optimization with SAP Transportation Management. International Journal of Engineering Research & Technology, 12(7).",
        "IJRISS (2025). Exploring the Influence of Kwacha Fluctuations on Medium Small and Micro Enterprises (MMSMEs) Profitability: A Case of Choma District, Southern Zambia. https://doi.org/10.47772/ijriss.2025.91100274",
        "IJIRMM. Effect of Exchange Rate on Pricing in Business Operations of SMEs in Enugu State. https://sadijournals.org/index.php/IJIRMM/article/view/643",
        "Lubanovic, B. (2019). Introducing Python (2nd ed.). O’Reilly Media.",
        "Mangan, J., & Lalwani, C. (2016). Global Logistics and Supply Chain Management. John Wiley & Sons.",
        "Nagle, T. T., & Müller, G. The Strategy and Tactics of Pricing. Pearson.",
        "NetSuite. Why You Need to Calculate Landed Cost. https://www.netsuite.com/portal/resource/articles/inventory-management/landed-cost.shtml",
        "Odoo Documentation — Inventory / Purchase / Landed Costs. https://www.odoo.com/documentation/",
        "Pressman, R. S., & Maxim, B. R. (2020). Software Engineering: A Practitioner’s Approach (9th ed.). McGraw-Hill.",
        "DiVA portal thesis: The Impact of Exchange Rate Volatility on Profitability and Hedging Behavior of Import Dependent Firms. https://www.diva-portal.org/smash/get/diva2:2012289/FULLTEXT01.pdf",
    ]
    for i, r in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(-1.25)
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        set_run_font(p.add_run(f"[{i}] {r}"), size=BODY_SIZE)


def appendix_a(doc, author):
    partner = partner_of(author)
    page_break(doc)
    add_heading_styled(doc, "Appendix A  Individual Contribution Statement", 1)
    add_para(
        doc,
        f"Purpose: identify work for which {author['name']} was primarily responsible, "
        "consistent with the main report, and acknowledge the co-author’s complementary work.",
        first_line=False,
        italic=True,
    )
    add_heading_styled(doc, "A.1 Project and Group Information", 2)
    add_table(
        doc,
        ["Item", "Details"],
        [
            ["Project Title", PROJECT_TITLE],
            ["Student Name", author["name"]],
            ["Computer Number", author["computer"]],
            ["Group Size", "2"],
            ["Co-author", f"{partner['name']} ({partner['computer']})"],
            ["Supervisor", SUPERVISOR],
            ["Source repository", REPO_URL],
            ["Submission form", "Individual report only (shared code via Git URL)"],
        ],
    )

    add_heading_styled(doc, "A.2 Overall Group Project", 2)
    add_para(
        doc,
        "The group developed Tranz-Algo, an explainable landed-cost inventory orchestration system "
        "with dual cost modes, True Cost studio, Cost history, and batch-linked stock for import SMEs.",
    )

    add_heading_styled(doc, "A.3 Individual Role and Responsibilities", 2)
    add_para(
        doc,
        f"Primary role: {author['role_short']}.",
        first_line=False,
    )
    add_para(doc, author["appendix_detail"], first_line=False)

    add_heading_styled(doc, "A.4 Collaboration and Co-author Work", 2)
    add_para(
        doc,
        "Joint (both): problem definition; architecture; technology selection and proposal deltas; "
        "ERD direction; dual modes and formulas; evaluation approach; presentation narrative.",
        first_line=False,
    )
    add_para(
        doc,
        f"Co-author ({partner['name']}): {partner['appendix_detail']}",
        first_line=False,
    )
    add_para(
        doc,
        "Estimated contribution (honest): design ownership roughly shared; implementation effort "
        "majority Nixon Shibuchi, with substantive UI/True Cost contribution from Musonda Katongo. "
        "Neither member claims the other’s exclusive code as sole personal work.",
        first_line=False,
    )
    add_para(doc, "Supervisor verification: ______________________  Date: __________", first_line=False)


def appendix_b(doc):
    page_break(doc)
    add_heading_styled(doc, "Appendix B  User Manual / Installation Guide (summary)", 1)
    add_para(
        doc,
        f"Full source, README, and demo CSVs are in the Git repository: {REPO_URL}. "
        "Operator and demo guidance is summarised below; CSV samples live under docs/scenarios/samples/.",
        first_line=False,
    )
    add_bullet(doc, "Clone: git clone https://github.com/Nixon-Shibuchi/Tranz-Algo.git")
    add_bullet(doc, "Backend: create venv, pip install -r requirements.txt, set DATABASE_URL, uvicorn app.main:app --reload")
    add_bullet(doc, "Frontend: npm install && npm run dev")
    add_bullet(doc, "Demo: Import products_demo.csv, then shipment_A_full.csv / shipment_B_full.csv from docs/scenarios/samples/")
    add_bullet(doc, "True Cost → allocate → price → commit → Cost history; optional Record sale")


def build_one(author):
    doc = Document()
    configure_print_styles(doc)

    # Section 1 — front matter: lower-case Roman numerals (cover has no footer number)
    front = doc.sections[0]
    configure_section_geometry(front)
    front.different_first_page_header_footer = True
    set_page_number_format(front, fmt="lowerRoman", start=1)
    add_page_number(front, prefix="")  # i, ii, iii …

    build_cover(doc, author)
    build_front_matter(doc, author)
    build_table_of_contents(doc)
    add_heading_styled(doc, "LIST OF ABBREVIATIONS AND ACRONYMS", 1)
    add_table(
        doc,
        ["Abbreviation", "Meaning"],
        [
            ["API", "Application Programming Interface"],
            ["CSV", "Comma-Separated Values"],
            ["ERP", "Enterprise Resource Planning"],
            ["FX", "Foreign Exchange"],
            ["JWT", "JSON Web Token"],
            ["MVP", "Minimum Viable Product"],
            ["POS", "Point of Sale"],
            ["SME", "Small and Medium-sized Enterprise"],
            ["UI", "User Interface"],
            ["VAT", "Value Added Tax"],
            ["ZMW", "Zambian Kwacha"],
        ],
    )

    # Section 2 — body from Chapter 1: Arabic numerals restarting at 1
    body = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section_geometry(body)
    body.different_first_page_header_footer = False
    set_page_number_format(body, fmt="decimal", start=1)
    add_page_number(body, prefix="")  # 1, 2, 3 …

    chapter_1(doc)
    page_break(doc)
    chapter_2(doc)
    page_break(doc)
    chapter_3(doc)
    page_break(doc)
    chapter_4(doc)
    page_break(doc)
    chapter_5(doc, author)
    page_break(doc)
    chapter_6(doc, author)
    page_break(doc)
    chapter_7(doc)
    page_break(doc)
    chapter_8(doc)
    page_break(doc)
    references(doc)
    appendix_a(doc, author)
    appendix_b(doc)

    out = OUT_DIR / f"Tranz-Algo_Final_Year_Project_Report_2026_{author['slug']}.docx"
    try:
        doc.save(out)
    except PermissionError:
        alt = OUT_DIR / f"Tranz-Algo_Final_Year_Project_Report_2026_{author['slug']}_updated.docx"
        doc.save(alt)
        print("LOCKED:", out)
        print(alt)
        return alt
    print(out)
    return out


def build():
    draw_all()
    paths = [build_one(AUTHORS["nixon"]), build_one(AUTHORS["musonda"])]
    return paths


if __name__ == "__main__":
    build()
