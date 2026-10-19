"""
Build a clean Tranz-Algo FYP presentation following the 2026 template slide order.
Sized for projector / printed handouts: large type, black body text, high-DPI diagrams.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from draw_figures import draw_all

OUT_DIR = Path(__file__).resolve().parent
OUT = OUT_DIR / "Tranz-Algo_Final_Year_Project_Presentation_2026.pptx"
FIG = OUT_DIR / "figures"
LOGO_UNZA = FIG / "unza_logo.png"
LOGO_TRANZ = FIG / "tranzalgo_logo.png"

# Widescreen 16:9 (same as UNZA 2026 template)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# High-contrast palette (matches print-oriented report figures)
NAVY = RGBColor(0x10, 0x2A, 0x38)
TEAL = RGBColor(0x1F, 0x5A, 0x68)
ACCENT = RGBColor(0xA8, 0x48, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE8, 0xEE, 0xF1)
MUTED = RGBColor(0x2A, 0x35, 0x3C)  # dark enough when projected
DARK = RGBColor(0x00, 0x00, 0x00)

# Visibility defaults for rooms / printouts
TITLE_SIZE = 32
BODY_SIZE = 20
CAPTION_SIZE = 16
FOOTER_SIZE = 13
TABLE_HEADER = 16
TABLE_BODY = 15
FONT = "Calibri"

STUDENT_1 = "Nixon Shibuchi"
STUDENT_1_COMP = "2021437361"
STUDENT_2 = "Musonda Katongo"
STUDENT_2_COMP = "2020583969"
SUPERVISOR = "Mrs. Leena Kumar"
REPO_URL = "https://github.com/Nixon-Shibuchi/Tranz-Algo"


def _run(p, text, *, size=BODY_SIZE, bold=False, color=DARK, font=FONT):
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return run


def _set_fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def style_table_cell(cell, *, size, bold=False, color=DARK, fill=None):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    for p in cell.text_frame.paragraphs:
        p.space_before = Pt(4)
        p.space_after = Pt(4)
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = FONT


def add_chrome(slide, page: int, total: int = 17):
    """Top accent bar + footer (readable when projected or printed)."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.16))
    _set_fill(bar, TEAL)

    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, SLIDE_H - Inches(0.48), SLIDE_W, Inches(0.48)
    )
    _set_fill(foot, LIGHT)
    tb = slide.shapes.add_textbox(
        Inches(0.4), SLIDE_H - Inches(0.42), Inches(10.5), Inches(0.36)
    )
    p = tb.text_frame.paragraphs[0]
    _run(
        p,
        "Tranz-Algo  ·  Final Year Project  ·  Academic Year 2026",
        size=FOOTER_SIZE,
        bold=True,
        color=MUTED,
    )
    num = slide.shapes.add_textbox(
        SLIDE_W - Inches(1.4), SLIDE_H - Inches(0.42), Inches(1.1), Inches(0.36)
    )
    p = num.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, f"{page} / {total}", size=FOOTER_SIZE, bold=True, color=NAVY)


def add_title(slide, text: str, top=Inches(0.32)):
    box = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.3), Inches(0.75))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _run(p, text, size=TITLE_SIZE, bold=True, color=NAVY)
    return box


def add_bullets(
    slide,
    lines,
    *,
    left=Inches(0.6),
    top=Inches(1.15),
    width=Inches(12),
    height=Inches(5.2),
    size=BODY_SIZE,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(12)
        if ": " in line and not line.startswith("http"):
            head, rest = line.split(": ", 1)
            _run(p, f"•  {head}: ", size=size, bold=True, color=NAVY)
            _run(p, rest, size=size, color=DARK)
        else:
            _run(p, f"•  {line}", size=size, color=DARK)
    return box


def add_caption(slide, text, *, top):
    box = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.3), Inches(0.45))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _run(p, text, size=CAPTION_SIZE, bold=True, color=MUTED)
    return box


def fit_picture(slide, path: Path, *, top, max_width, max_height, left=None):
    """Place image scaled to fit within max_width × max_height."""
    from PIL import Image

    with Image.open(path) as im:
        iw, ih = im.size
    aspect = iw / float(ih)
    max_w = float(max_width)
    max_h = float(max_height)
    if aspect > (max_w / max_h):
        w = max_w
        h = w / aspect
    else:
        h = max_h
        w = h * aspect
    if left is None:
        left = (float(SLIDE_W) - w) / 2
    else:
        left = float(left)
    return slide.shapes.add_picture(str(path), int(left), int(top), width=int(w), height=int(h))


def blank_slide(prs):
    for layout in prs.slide_layouts:
        if "blank" in (layout.name or "").lower():
            return prs.slides.add_slide(layout)
    return prs.slides.add_slide(prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)])


def build():
    draw_all()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ---- 1 Title ----
    s = blank_slide(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    _set_fill(bg, NAVY)
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.05), SLIDE_W, Inches(0.14))
    _set_fill(accent, ACCENT)

    logo_y = Inches(0.4)
    logo_h = Inches(1.05)
    if LOGO_UNZA.exists():
        s.shapes.add_picture(str(LOGO_UNZA), Inches(4.55), logo_y, height=logo_h)
    if LOGO_TRANZ.exists():
        s.shapes.add_picture(str(LOGO_TRANZ), Inches(7.35), logo_y, height=logo_h)

    t = s.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.7), Inches(2.0))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, "Tranz-Algo", size=48, bold=True, color=WHITE)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(14)
    _run(
        p2,
        "Explainable Landed-Cost Inventory Orchestration\nfor Import-Dependent SMEs",
        size=24,
        bold=True,
        color=WHITE,
    )

    meta = s.shapes.add_textbox(Inches(0.8), Inches(5.35), Inches(11.7), Inches(1.8))
    tf = meta.text_frame
    tf.word_wrap = True
    for i, line in enumerate(
        [
            f"{STUDENT_1}  ·  {STUDENT_2}",
            f"{STUDENT_1_COMP}  ·  {STUDENT_2_COMP}",
            "Final Year Project Presentation",
            "University of Zambia  ·  School of Natural and Applied Sciences",
            "Department of Computing and Informatics  ·  2026",
            REPO_URL,
        ]
    ):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(3)
        _run(p, line, size=16 if i < 5 else 14, bold=(i == 0), color=WHITE)

    # ---- 2 Team ----
    s = blank_slide(prs)
    add_chrome(s, 2)
    add_title(s, "Team Roles and Responsibilities")
    rows, cols = 3, 3
    table = s.shapes.add_table(rows, cols, Inches(0.7), Inches(1.35), Inches(11.9), Inches(2.5)).table
    headers = ["Member", "Computer No.", "Primary Role / Responsibilities"]
    data = [
        [STUDENT_1, STUDENT_1_COMP, "Implementation lead · costing engine · API · evaluation · docs"],
        [STUDENT_2, STUDENT_2_COMP, "Co-design · literature review · UI / True Cost contribution · demo"],
    ]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        style_table_cell(cell, size=TABLE_HEADER, bold=True, color=WHITE, fill=TEAL)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            fill = LIGHT if i % 2 == 1 else None
            style_table_cell(cell, size=TABLE_BODY, color=DARK, fill=fill)
    sup = s.shapes.add_textbox(Inches(0.7), Inches(4.15), Inches(11.9), Inches(0.5))
    p = sup.text_frame.paragraphs[0]
    _run(p, f"Supervisor: {SUPERVISOR}", size=18, bold=True, color=NAVY)
    add_bullets(
        s,
        [
            "Design & technology: discussed jointly before coding",
            "Most implementation: Nixon (including pending updates)",
            "UI / True Cost workflow commit + literature review: Musonda",
        ],
        top=Inches(4.75),
        height=Inches(1.7),
        size=18,
    )

    # ---- 3 What is Tranz-Algo? ----
    s = blank_slide(prs)
    add_chrome(s, 3)
    add_title(s, "What is Tranz-Algo?")
    add_bullets(
        s,
        [
            "A web admin app for import retailers — not a till/POS, not customs software, not a full ERP",
            "You create a shipment batch from supplier invoices plus shipping, duty and VAT paperwork",
            "It splits those shared costs onto each product → a clear true landed cost per unit",
            "It suggests a safe retail price (your margin, with a 10% safety floor above landed cost)",
            "When you commit: Cost history is saved and shelf stock is updated for that batch",
            "Two modes: Already-Local (amounts already in ZMW) or Auto FX (foreign prices × exchange rate)",
        ],
        size=19,
    )

    # ---- 4 Overview ----
    s = blank_slide(prs)
    add_chrome(s, 4)
    add_title(s, "Presentation Overview")
    add_bullets(
        s,
        [
            "What Tranz-Algo is (and is not)",
            "Background and Context",
            "System Analysis and Design (architecture, data, algorithms)",
            "Implementation / Development",
            "Testing, Results and Evaluation",
            "Discussion",
            "Conclusion",
            "System Demonstration",
        ],
        size=20,
    )

    # ---- 5 Background ----
    s = blank_slide(prs)
    add_chrome(s, 5)
    add_title(s, "Background and Context")
    add_bullets(
        s,
        [
            "Import SMEs know invoices and clearing papers, yet struggle to recover per-item true cost",
            "Shared freight, duty and VAT plus mixed foreign/local paperwork erase margins",
            "FX movement can wipe cost-plus pricing after a selling price is set",
            "Users: retail administrators (desk workflow) — not till operators",
            "Tranz-Algo: landed cost → suggested retail → Cost history → batch stock",
        ],
        size=20,
    )

    # ---- 6 Architecture ----
    s = blank_slide(prs)
    add_chrome(s, 6)
    add_title(s, "System Architecture")
    add_caption(
        s,
        "React screens → FastAPI → services (costing / workflow / FX) → PostgreSQL · optional live FX · JWT",
        top=Inches(1.05),
    )
    fit_picture(
        s,
        FIG / "architecture.png",
        top=Inches(1.55),
        max_width=Inches(12.4),
        max_height=Inches(5.15),
    )

    # ---- 7 ERD ----
    s = blank_slide(prs)
    add_chrome(s, 7)
    add_title(s, "Database / Data Design")
    add_caption(
        s,
        "UML-style entities · duty/VAT as paperwork totals on the batch · Cost history from ledger · Sale = stock-out only",
        top=Inches(1.05),
    )
    fit_picture(
        s,
        FIG / "erd.png",
        top=Inches(1.5),
        max_width=Inches(12.5),
        max_height=Inches(5.25),
    )

    # ---- 8 Algorithms + dual modes ----
    s = blank_slide(prs)
    add_chrome(s, 8)
    add_title(s, "Algorithms / Models / Technical Design")
    add_bullets(
        s,
        [
            "Auto FX: foreign unit × exchange rate",
            "Already-Local: local unit product cost (no Fetch FX required)",
            "Allocate shipping / duty / VAT by quantity or weight; compare methods",
            "Suggested RP = Landed × (1 + margin)",
            "Safety floor = Landed × 1.10 (live on True Cost adjust)",
        ],
        left=Inches(0.45),
        top=Inches(1.15),
        width=Inches(6.5),
        height=Inches(5.3),
        size=17,
    )
    fit_picture(
        s,
        FIG / "dual_modes.png",
        top=Inches(1.4),
        max_width=Inches(5.8),
        max_height=Inches(4.8),
        left=Inches(7.15),
    )

    # ---- 9 Dev environment ----
    s = blank_slide(prs)
    add_chrome(s, 9)
    add_title(s, "Development Environment")
    add_bullets(
        s,
        [
            "Languages: Python 3, JavaScript (React)",
            "Frameworks: FastAPI, Vite/React, SQLAlchemy 2, Pydantic",
            "Database: PostgreSQL (primary), SQLite fallback",
            "FX: Frankfurter / open.er-api + cache / manual",
            "Tools: VS Code / Cursor, Git / GitHub",
            f"Repository: {REPO_URL}",
        ],
        size=19,
    )

    # ---- 10 Modules ----
    s = blank_slide(prs)
    add_chrome(s, 10)
    add_title(s, "System Components / Modules")
    add_bullets(
        s,
        [
            "Auth: JWT login / Create account",
            "Products: catalogue + CSV import (+ optional auto SKU)",
            "Shipments: dual modes, paperwork costs, CSV, Clear drafts",
            "True Cost studio: Lines → Allocate → Price → Commit",
            "FX service: Fetch / cache / starter 20 / manual",
            "Stock board + Record sale (stock-out only — not POS)",
            "Dashboard Start fresh + Cost history (When / What)",
        ],
        size=19,
    )

    # ---- 11 Key implementation + workflow ----
    s = blank_slide(prs)
    add_chrome(s, 11)
    add_title(s, "Key Implementation Details")
    add_bullets(
        s,
        [
            "React web admin instead of React Native — prioritise desk costing MVP",
            "Strategy (modes/allocation) · Facade (commit) · Chain (FX providers)",
            "CSV import drives Demo A/B: SHP-DEMO-LOCAL-ZMW / SHP-DEMO-AUTO-FX",
            "Security: bcrypt · JWT Bearer · Pydantic validation",
        ],
        top=Inches(1.1),
        height=Inches(2.5),
        size=18,
    )
    fit_picture(
        s,
        FIG / "true_cost_workflow.png",
        top=Inches(3.65),
        max_width=Inches(12.4),
        max_height=Inches(2.9),
    )

    # ---- 12 Challenges ----
    s = blank_slide(prs)
    add_chrome(s, 12)
    add_title(s, "Challenges and Solutions")
    add_bullets(
        s,
        [
            "Mixed foreign/local paperwork → Dual cost modes at shipment create",
            "FX API downtime → Cache, starter rate 20, manual override",
        ],
        size=20,
    )

    # ---- 13 Tests ----
    s = blank_slide(prs)
    add_chrome(s, 13)
    add_title(s, "Test Results")
    rows, cols = 7, 3
    table = s.shapes.add_table(rows, cols, Inches(0.55), Inches(1.2), Inches(12.2), Inches(5.3)).table
    headers = ["ID", "Test case", "Status"]
    data = [
        ["TC-01", "Already-Local Demo A vs expected spreadsheet costs (±0.01 ZMW)", "PASS"],
        ["TC-02", "Auto FX Demo B vs expected costs (rate 20)", "PASS"],
        ["TC-03", "Allocation compare qty vs weight (Demo A)", "PASS"],
        ["TC-04", "Safety floor on True Cost adjust", "PASS"],
        ["TC-05", "CSV UI demo path (Scenario 05)", "PASS"],
        ["TC-06", "Record sale stock-out (not POS)", "PASS"],
    ]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        style_table_cell(cell, size=TABLE_HEADER, bold=True, color=WHITE, fill=TEAL)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            fill = LIGHT if i % 2 == 1 else None
            color = TEAL if j == 2 else DARK
            style_table_cell(cell, size=TABLE_BODY, bold=(j == 2), color=color, fill=fill)

    # ---- 14 Discussion ----
    s = blank_slide(prs)
    add_chrome(s, 14)
    add_title(s, "Discussion of Findings")
    add_bullets(
        s,
        [
            "Demo A/B costs matched our pre-calculated spreadsheet answers — formulas are correct",
            "Dual modes address mixed paperwork — the original problem",
            "Design, build and evaluate objectives were met for MVP scope",
            "Key finding: explainable batch costing can be accurate and usable",
        ],
        size=20,
    )

    # ---- 15 Limitations ----
    s = blank_slide(prs)
    add_chrome(s, 15)
    add_title(s, "Limitations")
    add_bullets(
        s,
        [
            "Does not legally assess customs duty/VAT — paperwork totals only",
            "Auto FX / Demo B number-check use rate 20 (same as UI starter)",
            "Value-based allocation is stretch in the UI",
            "No large-scale SME field trial yet — evaluation is spreadsheet checks + live demo",
        ],
        size=20,
    )

    # ---- 16 Conclusion ----
    s = blank_slide(prs)
    add_chrome(s, 16)
    add_title(s, "Conclusion")
    add_bullets(
        s,
        [
            "Tranz-Algo delivers explainable import true cost, suggested retail, Cost history and stock",
            "Design was joint; implementation mostly Nixon with Musonda’s UI / True Cost contribution",
            "Next: optional mobile companion, cloud Postgres, usability study",
            "Thank you — demonstration follows",
        ],
        size=20,
    )

    # ---- 17 Questions ----
    s = blank_slide(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    _set_fill(bg, NAVY)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(3.45), SLIDE_W, Inches(0.12))
    _set_fill(bar, ACCENT)
    t = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.2))
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, "Questions and Discussion", size=40, bold=True, color=WHITE)
    t2 = s.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(11.7), Inches(1.6))
    p = t2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, "Thank you", size=28, bold=True, color=WHITE)
    p2 = t2.text_frame.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(10)
    _run(p2, "Tranz-Algo  ·  University of Zambia  ·  2026", size=18, color=WHITE)
    p3 = t2.text_frame.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(8)
    _run(p3, REPO_URL, size=16, color=WHITE)

    try:
        prs.save(str(OUT))
        print(OUT)
        return OUT
    except PermissionError:
        alt = OUT.with_name(OUT.stem + "_updated.pptx")
        prs.save(str(alt))
        print(alt)
        return alt


if __name__ == "__main__":
    build()
